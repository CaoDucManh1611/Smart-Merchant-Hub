from app.rag.semantic_chunker import chunk_document


def test_catalog_rows_never_mix_products():
    text = "sku | name | price\nA-001 | Bình A | 100000\nB-002 | Bình B | 200000"
    chunks = chunk_document(text, "products.csv", chunk_size=35)
    assert len(chunks) == 2
    assert "Bình A" in chunks[0].content and "Bình B" not in chunks[0].content
    assert "Bình B" in chunks[1].content and "Bình A" not in chunks[1].content
    assert chunks[0].metadata["parent_id"] != chunks[1].metadata["parent_id"]


def test_policy_keeps_condition_and_exception_in_one_parent():
    text = "# Đổi trả\nKhách được đổi hàng trong 7 ngày.\n\nChỉ áp dụng khi còn tem.\nKhông áp dụng với hàng đã sử dụng.\n\n# Bảo hành\nBảo hành 12 tháng."
    chunks = chunk_document(text, "policy.md", chunk_size=45)
    policy = [chunk for chunk in chunks if chunk.metadata["section"] == "Đổi trả"]
    assert policy
    assert "còn tem" in " ".join(chunk.content for chunk in policy)
    assert "Không áp dụng" in " ".join(chunk.content for chunk in policy)
    assert all(chunk.metadata["parent_id"] == policy[0].metadata["parent_id"] for chunk in policy)
    assert all(chunk.metadata["parent_id"] != policy[0].metadata["parent_id"] for chunk in chunks if chunk.metadata["section"] == "Bảo hành")


def test_gemini_boundaries_only_group_original_paragraphs(monkeypatch):
    from app.rag import semantic_chunker

    monkeypatch.setattr(semantic_chunker, "_semantic_boundaries", lambda parts: [2, 3])
    text = "Câu đầu.\n\nĐiều kiện áp dụng.\n\nNgoại lệ."
    chunks = chunk_document(text, "notes.txt", chunk_size=20)
    assert chunks[0].content == "Câu đầu.\n\nĐiều kiện áp dụng."
    assert chunks[1].content == "Ngoại lệ."


def test_bad_model_boundary_falls_back_without_losing_source(monkeypatch):
    from app.rag import semantic_chunker

    monkeypatch.setattr(semantic_chunker, "_semantic_boundaries", lambda parts: [99])
    text = "Một.\n\nHai."
    chunks = chunk_document(text, "notes.txt", chunk_size=10)
    assert "Một." in " ".join(chunk.content for chunk in chunks)
    assert "Hai." in " ".join(chunk.content for chunk in chunks)


def test_neighbor_expansion_never_crosses_parent_or_shop():
    from types import SimpleNamespace
    from app.rag.retriever import RetrievedChunk, _expand_parent_neighbors

    lead = RetrievedChunk(2, 10, "Quy định đổi trả", 0.9, {"parent_id": 1}, "policy.md", 1)

    class Session:
        def execute(self, statement, params):
            assert params["business_id"] == 7
            assert params["document_id"] == 10
            assert "d.business_id = :business_id" in str(statement)
            return SimpleNamespace(fetchall=lambda: [
                SimpleNamespace(_mapping={"id": 1, "document_id": 10, "content": "Điều kiện còn tem", "chunk_metadata": {"parent_id": 1}, "source_filename": "policy.md", "chunk_index": 0}),
                SimpleNamespace(_mapping={"id": 3, "document_id": 10, "content": "Giá sản phẩm khác", "chunk_metadata": {"parent_id": 2}, "source_filename": "policy.md", "chunk_index": 2}),
            ])

    expanded = _expand_parent_neighbors([lead], Session(), 7, 3)
    assert [item.chunk_id for item in expanded] == [2, 1]


def test_catalog_lines_below_one_heading_remain_separate_products():
    text = "# Danh sách sản phẩm\n- Bình A (SKU A-001): giá 100000, tồn kho: 2\n- Bình B (SKU B-002): giá 200000, tồn kho: 3"
    chunks = chunk_document(text, "products.md")
    products = [chunk for chunk in chunks if chunk.metadata["unit_type"] == "product"]
    assert len(products) == 2
    assert products[0].metadata["sku"] == "A-001"
    assert products[1].metadata["sku"] == "B-002"
    assert products[0].metadata["parent_id"] != products[1].metadata["parent_id"]


def test_gemini_groups_original_policy_paragraphs_without_rewriting():
    calls = []

    def detect(parts):
        calls.append(parts)
        return [3, 4]

    text = "# Đổi trả\nĐược đổi trong 7 ngày.\n\nChỉ khi còn tem.\n\nKhông áp dụng với hàng đã dùng.\n\nLiên hệ nhân viên để được hỗ trợ."
    chunks = chunk_document(text, "policy.md", semantic_boundary_detector=detect)
    assert calls == [["# Đổi trả\nĐược đổi trong 7 ngày.", "Chỉ khi còn tem.", "Không áp dụng với hàng đã dùng.", "Liên hệ nhân viên để được hỗ trợ."]]
    assert [chunk.content for chunk in chunks] == [
        "# Đổi trả\nĐược đổi trong 7 ngày.\n\nChỉ khi còn tem.\n\nKhông áp dụng với hàng đã dùng.",
        "Liên hệ nhân viên để được hỗ trợ.",
    ]
    assert len({chunk.metadata["parent_id"] for chunk in chunks}) == 1


def test_invalid_or_failed_gemini_boundaries_preserve_original_information():
    text = "Quy tắc chính.\n\nChỉ áp dụng khi còn tem.\n\nNgoại lệ: hàng đã dùng."
    for detector in (lambda _: [2, 2, 3], lambda _: [9], lambda _: (_ for _ in ()).throw(RuntimeError("private token"))):
        chunks = chunk_document(text, "policy.txt", semantic_boundary_detector=detector)
        combined = "\n\n".join(chunk.content for chunk in chunks)
        assert all(part in combined for part in ("Quy tắc chính.", "Chỉ áp dụng khi còn tem.", "Ngoại lệ: hàng đã dùng."))
        assert "private token" not in combined


def test_semantic_chunking_batches_long_documents_and_never_sends_csv(monkeypatch):
    from app.rag import semantic_chunker

    batches = []

    def detect(parts):
        batches.append(parts)
        return [len(parts)]

    text = "\n\n".join(f"Đoạn {index}: chính sách của shop." for index in range(45))
    chunks = chunk_document(text, "policy.txt", semantic_boundary_detector=detect)
    assert len(batches) > 1
    assert all(len(batch) <= semantic_chunker._MAX_SEMANTIC_PARAGRAPHS for batch in batches)
    assert all(f"Đoạn {index}:" in " ".join(chunk.content for chunk in chunks) for index in range(45))

    chunk_document("sku | name\nA-001 | Bình A", "products.csv", semantic_boundary_detector=lambda _: (_ for _ in ()).throw(AssertionError("CSV should not call Gemini")))


def test_gemini_reply_must_be_exact_json_boundaries():
    from app.rag.semantic_chunker import _parse_model_boundaries

    assert _parse_model_boundaries("[2, 4]", 4) == [2, 4]
    for value in ('```json\n[2, 4]\n```', '[2, 2, 4]', '[4, 3]', '[1, 5]', '{"ends":[4]}', '[true, 4]'):
        assert _parse_model_boundaries(value, 4) is None


def test_gemini_chunking_uses_dedicated_key_and_preserves_original_payload(monkeypatch):
    from app.rag import llm_caller

    monkeypatch.setattr(llm_caller.settings, "GEMINI_API_KEY", "test-only-key")
    monkeypatch.setattr(llm_caller.settings, "GEMINI_API_KEYS", "")
    captured = []

    def fake_call(messages, key, model):
        captured.append((messages, key, model))
        return "[2]"

    monkeypatch.setattr(llm_caller, "_call_gemini_once", fake_call)
    messages = llm_caller.gemini_chunking_messages(["Quy định.", "Ngoại lệ."])
    assert "[1] Quy định." in messages[1]["content"]
    assert "[2] Ngoại lệ." in messages[1]["content"]
    assert llm_caller.call_gemini_for_chunking(messages) == "[2]"
    assert captured[0][1] == "test-only-key"


def test_freeform_model_chunks_share_one_parent():
    chunks = chunk_document("Quy định.\n\nNgoại lệ.", "policy.txt",
                            semantic_boundary_detector=lambda parts: [1, 2])
    assert len(chunks) == 2
    assert chunks[0].metadata["parent_id"] == chunks[1].metadata["parent_id"]


def test_long_paragraph_splits_at_complete_sentences_and_preserves_decimal():
    text = "Giá là 1.25 USD. " + "Sản phẩm được bảo hành đầy đủ trong mười hai tháng. " * 20
    seen = []
    def detector(parts):
        seen.extend(parts)
        return list(range(1, len(parts) + 1))
    chunks = chunk_document(text, "policy.txt", chunk_size=120, semantic_boundary_detector=detector)
    assert len(chunks) > 1
    assert any("1.25 USD." in part for part in seen)
    assert " ".join(" ".join(c.content for c in chunks).split()) == " ".join(text.split())
    assert all(c.content.rstrip().endswith(".") for c in chunks)


def test_contextual_embedding_uses_heading_path_but_keeps_citation_original():
    from app.rag.semantic_chunker import contextual_embedding_text
    chunks = chunk_document("# Chính sách\n## Đổi trả\nKhách được đổi trong 7 ngày.", "shop.md")
    child = chunks[-1]
    assert child.metadata["section"] == "Chính sách > Đổi trả"
    assert "shop.md" in contextual_embedding_text(child)
    assert "Chính sách > Đổi trả" in contextual_embedding_text(child)
    assert contextual_embedding_text(child).endswith(child.content)
    assert not child.content.startswith("Document:")
    assert child.metadata["parent_start"] <= child.index <= child.metadata["parent_end"]


def test_markdown_table_repeats_header_without_mixing_rows():
    text = "# Hàng hóa\n| SKU | Giá |\n| --- | --- |\n| A-1 | 100 |\n| B-2 | 200 |"
    rows = [c for c in chunk_document(text, "catalog.md") if c.metadata["unit_type"] == "table_row"]
    assert len(rows) == 2
    assert "| SKU | Giá |" in rows[0].content
    assert "B-2" not in rows[0].content


def test_complete_parent_retrieval_keeps_source_order_and_other_hits():
    from types import SimpleNamespace
    from app.rag.retriever import RetrievedChunk, _expand_parent_neighbors
    lead = RetrievedChunk(3, 10, "Ngoại lệ.", .9, {"parent_id": 1, "parent_start": 0, "parent_end": 2}, "policy.md", 2)
    other = RetrievedChunk(9, 20, "Bảo hành.", .8, {}, "warranty.md", 0)
    class Session:
        def execute(self, statement, params):
            assert params["business_id"] == 7
            assert params["document_id"] == 10
            assert "d.business_id = :business_id" in str(statement)
            return SimpleNamespace(fetchall=lambda: [SimpleNamespace(_mapping={
                "id": i+1, "document_id": 10, "content": content,
                "chunk_metadata": {"parent_id": 1}, "source_filename": "policy.md", "chunk_index": i,
            }) for i, content in enumerate(["Quy định.", "Điều kiện.", "Ngoại lệ."])])
    results = _expand_parent_neighbors([lead, other], Session(), 7, 2)
    assert results[0].content == "Quy định.\n\nĐiều kiện.\n\nNgoại lệ."
    assert results[0].metadata["expanded_chunk_ids"] == [1, 2, 3]
    assert results[1].chunk_id == 9


def test_table_without_heading_and_multiline_product_keep_details():
    rows = chunk_document("| SKU | Giá |\n| --- | --- |\n| A-1 | 100 |", "table.md")
    assert len(rows) == 1 and rows[0].metadata["unit_type"] == "table_row"
    chunks = chunk_document("# Hàng\n- Bình A (SKU A-1)\n  Giá 100.\n- Bình B (SKU B-2)\n  Giá 200.", "catalog.md")
    products = [c for c in chunks if c.metadata["unit_type"] == "product"]
    assert "Giá 100." in products[0].content and "200" not in products[0].content
    assert "Giá 200." in products[1].content


def test_ingestion_embeds_context_and_stores_only_source(monkeypatch):
    from types import SimpleNamespace
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from app.models import Business, Document, DocumentChunk
    from app.services import ingestion_service
    engine = create_engine("sqlite://")
    Business.metadata.create_all(engine)
    captured = []
    def embed(texts, **kwargs):
        captured.extend(texts)
        return [[0.0] * 768 for _ in texts]
    monkeypatch.setattr(ingestion_service, "embed_texts", embed)
    monkeypatch.setattr(ingestion_service.settings, "EMBEDDING_PROVIDER", "local")
    monkeypatch.setattr(ingestion_service.settings, "GEMINI_API_KEY", "")
    monkeypatch.setattr(ingestion_service.settings, "GEMINI_API_KEYS", "")
    monkeypatch.setattr(ingestion_service, "check_quota", lambda *a, **k: SimpleNamespace(allowed=True))
    monkeypatch.setattr(ingestion_service, "sync_catalog_products", lambda *a, **k: 0)
    with Session(engine) as db:
        business = Business(name="Test", slug="context-test")
        db.add(business)
        db.flush()
        doc = Document(business_id=business.id, filename="policy.md", file_type="md", file_size=0)
        db.add(doc)
        db.commit()
        ingestion_service.ingest_document(doc.id, "# Đổi trả\nKhách được đổi trong 7 ngày.".encode(), doc.filename, db, business_id=business.id)
        chunk = db.query(DocumentChunk).filter_by(document_id=doc.id).one()
        assert doc.status == "ready"
        assert captured[0].startswith("Document: policy.md\nSection: Đổi trả")
        assert chunk.content == "# Đổi trả\nKhách được đổi trong 7 ngày."
        assert chunk.metadata_["chunking_version"] == "structure-semantic-context-v2"


def test_html_and_docx_headings_survive_extraction():
    import io
    from docx import Document
    from app.rag.loader import load_document
    doc = Document()
    doc.add_heading("Chính sách", level=1)
    doc.add_heading("Đổi trả", level=2)
    doc.add_paragraph("Được đổi trong 7 ngày.")
    data = io.BytesIO()
    doc.save(data)
    for raw, name in [(data.getvalue(), "policy.docx"), (b"<h1>Policy</h1><h2>Returns</h2><p>Seven days.</p>", "policy.html")]:
        text = load_document(raw, name)
        assert "## " in text
        chunks = chunk_document(text, name)
        assert " > " in chunks[-1].metadata["section"]


def test_csv_multiline_cell_keeps_one_product_row():
    from app.rag.loader import load_document
    raw = b'sku,name,description\nA-1,Bottle,"Line one\nLine two"\nB-2,Cup,Small'
    chunks = chunk_document(load_document(raw, "products.csv"), "products.csv")
    assert len(chunks) == 2
    assert "Line one" in chunks[0].content and "Line two" in chunks[0].content
    assert "B-2" not in chunks[0].content


def test_large_parent_uses_bounded_neighbors_and_rejects_foreign_evidence():
    from types import SimpleNamespace
    from app.rag.retriever import RetrievedChunk, _expand_parent_neighbors
    lead = RetrievedChunk(11, 10, "Quy định.", .9, {"parent_id": 1, "parent_start": 0, "parent_end": 30}, "policy.md", 10)
    class Session:
        def execute(self, statement, params):
            assert params["start_index"] == 9 and params["end_index"] == 11
            rows = [
                (10, 10, 9, 1, "Điều kiện."),
                (12, 10, 11, 1, "x" * 7000),
                (99, 99, 9, 1, "Other document secret"),
                (98, 10, 9, 2, "Other parent secret"),
            ]
            return SimpleNamespace(fetchall=lambda: [SimpleNamespace(_mapping={
                "id": id, "document_id": doc, "chunk_index": index, "chunk_metadata": {"parent_id": parent},
                "content": content, "source_filename": "policy.md",
            }) for id, doc, index, parent, content in rows])
    result = _expand_parent_neighbors([lead], Session(), 7, 1)[0]
    assert result.content == "Điều kiện.\n\nQuy định."
    assert result.metadata["retrieval_context"] == "neighbors"
    assert result.metadata["expanded_chunk_ids"] == [10, 11]
